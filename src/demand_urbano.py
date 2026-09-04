"""Construye un punto de demanda urbano sintetico por distrito.

Las capas publicas de centros poblados de INEI/IGN solo cubren poblacion
*dispersa* (rural). La poblacion urbana concentrada no existe como puntos
en ninguna capa publica del geoportal, asi que se estima por diferencia:

    pob_urbana_estimada(distrito) = pob_total_proyectada(distrito) - suma(pob_total_dispersa en el distrito)

y se ubica en el punto "Capital de Distrito" del gazetteer nacional IGN
(con centroide del poligono distrital como respaldo si el distrito no
tiene un punto de capital en el gazetteer). Ver config.md, seccion
"Fuentes de datos y sustituciones documentadas", para el detalle
metodologico.

Entrada: archivos ya descargados por acquisition.py en data/raw/.
Salida: data/processed/demand_points_urbano.geojson
"""

import json
from collections import defaultdict

import openpyxl

from config import ROOT, load_config


def _load_district_population(xlsx_path, valid_ubigeos: set[str], anio: int) -> dict[str, int]:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))

    header_years = rows[4]  # fila: [None, None, 2018, 2019, ..., 2026]
    col_year = header_years.index(anio)

    poblacion: dict[str, int] = {}
    for row in rows[6:]:
        ubigeo = row[0]
        if not ubigeo or not isinstance(ubigeo, str):
            continue
        ubigeo = ubigeo.strip()
        if ubigeo not in valid_ubigeos:
            continue
        valor = row[col_year]
        if valor is not None:
            poblacion[ubigeo] = int(valor)
    return poblacion


def _load_valid_districts(admin3_path, departamentos: list[str]) -> list[dict]:
    dep_upper = {d.upper() for d in departamentos}
    data = json.loads(admin3_path.read_text(encoding="utf-8"))
    districts = []
    for feat in data["features"]:
        props = feat["properties"]
        if props["adm1_name"].upper() not in dep_upper:
            continue
        ubigeo = props["adm3_pcode"][2:]  # 'PE030101' -> '030101'
        districts.append(
            {
                "ubigeo": ubigeo,
                "dep": props["adm1_name"],
                "prov": props["adm2_name"],
                "dist": props["adm3_name"],
                "centroid_lon": props["center_lon"],
                "centroid_lat": props["center_lat"],
            }
        )
    return districts


def _load_dispersed_population_by_district(dest_dir, departamentos: list[str]) -> dict[str, int]:
    totals: dict[str, int] = defaultdict(int)
    for dep in departamentos:
        path = dest_dir / f"centros_poblados_dispersos_{dep.lower()}.geojson"
        data = json.loads(path.read_text(encoding="utf-8"))
        for feat in data["features"]:
            props = feat["properties"]
            cod_dist = props.get("cod_dist")
            pob = props.get("pob_total") or 0
            if cod_dist:
                totals[cod_dist] += pob
    return totals


def _find_codigo_field(props: dict) -> str | None:
    """El campo 'CODIGO' (centro poblado, 10 digitos: dep+prov+dist+secuencia)
    llega con el nombre corrupto desde el geoportal (perdida de byte en la
    tilde de 'CODIGO', no es un problema de nuestro lado). Se detecta por
    el patron del valor (10 digitos) en vez de por el nombre de columna."""
    for key, value in props.items():
        if isinstance(value, str) and len(value) == 10 and value.isdigit():
            return key
    return None


def _load_capital_points_by_district(dest_dir, departamentos: list[str]) -> dict[str, tuple[float, float]]:
    points: dict[str, tuple[float, float]] = {}
    for dep in departamentos:
        path = dest_dir / f"capitales_distritales_{dep.lower()}.geojson"
        data = json.loads(path.read_text(encoding="utf-8"))
        for feat in data["features"]:
            props = feat["properties"]
            codigo_field = _find_codigo_field(props)
            if codigo_field is None:
                continue
            ubigeo = props[codigo_field][:6]
            lon, lat = feat["geometry"]["coordinates"]
            points.setdefault(ubigeo, (lon, lat))
    return points


def build_demand_points_urbano(force: bool = False):
    cfg = load_config()
    departamentos = list(cfg["departamentos"].values())
    anio = cfg["fuentes"]["poblacion_distrital_anio"]

    raw_boundaries = ROOT / "data" / "raw" / "boundaries" / "per_admin3.geojson"
    raw_centros_poblados = ROOT / "data" / "raw" / "centros_poblados"
    raw_poblacion = ROOT / "data" / "raw" / "poblacion" / "poblacion_distrital_proyectada_2018_2026.xlsx"
    dest = ROOT / "data" / "processed" / "demand_points_urbano.geojson"
    dest.parent.mkdir(parents=True, exist_ok=True)

    if dest.exists() and not force:
        print(f"[demand_urbano] {dest.name} ya existe, se omite (force=True para regenerar).")
        return dest

    districts = _load_valid_districts(raw_boundaries, departamentos)
    valid_ubigeos = {d["ubigeo"] for d in districts}
    poblacion_total = _load_district_population(raw_poblacion, valid_ubigeos, anio)
    poblacion_dispersa = _load_dispersed_population_by_district(raw_centros_poblados, departamentos)
    capitales = _load_capital_points_by_district(raw_centros_poblados, departamentos)

    features = []
    sin_poblacion = []
    usando_centroide = []
    for d in districts:
        ubigeo = d["ubigeo"]
        total = poblacion_total.get(ubigeo)
        if total is None:
            sin_poblacion.append(ubigeo)
            continue
        dispersa = poblacion_dispersa.get(ubigeo, 0)
        urbana_estimada = max(total - dispersa, 0)

        if ubigeo in capitales:
            lon, lat = capitales[ubigeo]
            point_source = "capital_distrital"
        else:
            lon, lat = d["centroid_lon"], d["centroid_lat"]
            point_source = "centroide_distrito"
            usando_centroide.append(ubigeo)

        features.append(
            {
                "type": "Feature",
                "geometry": {"type": "Point", "coordinates": [lon, lat]},
                "properties": {
                    "ubigeo": ubigeo,
                    "dep": d["dep"],
                    "prov": d["prov"],
                    "dist": d["dist"],
                    "pob_total_proyectada": total,
                    "pob_dispersa": dispersa,
                    "pob_urbana_estimada": urbana_estimada,
                    "point_source": point_source,
                },
            }
        )

    geojson = {"type": "FeatureCollection", "features": features}
    dest.write_text(json.dumps(geojson, ensure_ascii=False), encoding="utf-8")

    print(f"[demand_urbano] {len(features)} distritos procesados -> {dest}")
    if sin_poblacion:
        print(f"[demand_urbano] AVISO: {len(sin_poblacion)} distritos sin poblacion proyectada (ubigeo no encontrado en el excel): {sin_poblacion}")
    if usando_centroide:
        print(f"[demand_urbano] AVISO: {len(usando_centroide)} distritos sin punto 'Capital de Distrito', se uso el centroide del poligono: {usando_centroide}")

    dept_totals_excel = _load_department_totals(raw_poblacion, departamentos)
    for dep in departamentos:
        dep_feats = [f["properties"] for f in features if f["properties"]["dep"].upper() == dep.upper()]
        total = sum(f["pob_total_proyectada"] for f in dep_feats)
        urbana = sum(f["pob_urbana_estimada"] for f in dep_feats)
        dispersa = sum(f["pob_dispersa"] for f in dep_feats)
        dept_total_excel = dept_totals_excel.get(dep.upper())
        gap = dept_total_excel - total if dept_total_excel is not None else None
        print(
            f"[demand_urbano] {dep}: {len(dep_feats)} distritos, pob. total {total}, urbana est. {urbana}, "
            f"dispersa {dispersa} | fila departamento (excel) {dept_total_excel}, diferencia {gap}"
        )
        if gap:
            print(
                f"[demand_urbano]   -> revisar: {dep} tiene {gap} hab. en la fila de departamento que no "
                f"caen en ningun distrito de nuestros limites administrativos (distritos nuevos no reflejados "
                f"en la capa de limites 2020, o ajuste propio de INEI en su tabla). Ver config.md."
            )

    return dest


def _load_department_totals(xlsx_path, departamentos: list[str]) -> dict[str, int]:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    col_year = rows[4].index(load_config()["fuentes"]["poblacion_distrital_anio"])
    dep_upper = {d.upper() for d in departamentos}
    totals = {}
    for row in rows[6:]:
        ubigeo, nombre = row[0], row[1]
        if isinstance(ubigeo, str) and ubigeo.endswith("0000") and isinstance(nombre, str) and nombre.upper() in dep_upper:
            totals[nombre.upper()] = row[col_year]
    return totals


if __name__ == "__main__":
    build_demand_points_urbano()
